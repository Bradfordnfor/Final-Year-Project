import csv
import io
from typing import Optional


def parse_int(raw, default: int) -> int:
    """Best-effort int parse. Blank/None/unparseable returns ``default``.

    Accepts spreadsheet-style floats like "2.0" (Excel stores numbers as floats).
    """
    if raw is None:
        return default
    text = str(raw).strip()
    if not text:
        return default
    try:
        return int(float(text))
    except (ValueError, TypeError):
        return default


def _normalize_cell(value) -> str:
    return "" if value is None else str(value).strip()


def read_rows(filename: str, content: bytes) -> tuple[list[str], list[dict]]:
    """Read a CSV or .xlsx upload into (header, rows).

    ``header`` is the lower-cased, trimmed column names. Each row is a dict of
    lower-cased column name -> stripped string value. Raises ``ValueError`` when
    the bytes cannot be read in the format implied by the filename.
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _read_xlsx(content)
    return _read_csv(content)


def _read_csv(content: bytes) -> tuple[list[str], list[dict]]:
    try:
        text = content.decode("utf-8-sig")
    except (UnicodeDecodeError, ValueError):
        raise ValueError("file is not decodable text")
    reader = csv.DictReader(io.StringIO(text))
    header = [(_normalize_cell(h)).lower() for h in (reader.fieldnames or [])]
    rows = []
    for raw in reader:
        rows.append({
            (_normalize_cell(k)).lower(): _normalize_cell(v)
            for k, v in raw.items() if k is not None
        })
    return header, rows


def _read_xlsx(content: bytes) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("file is not a readable .xlsx")
    ws = wb.active
    if ws is None:
        return [], []
    row_iter = ws.iter_rows(values_only=True)
    try:
        header_cells = next(row_iter)
    except StopIteration:
        return [], []
    header = [(_normalize_cell(c)).lower() for c in header_cells]
    rows = []
    for cells in row_iter:
        if cells is None:
            continue
        if all(_normalize_cell(c) == "" for c in cells):
            continue
        row = {}
        for i, key in enumerate(header):
            if not key:
                continue
            row[key] = _normalize_cell(cells[i]) if i < len(cells) else ""
        rows.append(row)
    return header, rows
