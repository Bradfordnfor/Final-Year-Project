r"""Apply this session's de-AI text edits (dash removal + sentence rewrites)
directly into the hand-edited Dissertation.docx, preserving the cover, the 22
figures, the table of contents, and all inline bold/italic formatting.

Method (formatting-preserving):
  * Build a map  normalise(OLD paragraph text) -> NEW paragraph text  by diffing
    the pre-de-AI markdown (git b34c947) against the current markdown.
  * For each paragraph (and table cell) in the docx, if its normalised text
    matches an OLD entry, rebuild it from the NEW text, re-anchoring the
    paragraph's existing bold/italic runs (the technical terms, which never
    changed) so their formatting is kept exactly.
  * Normalisation strips ALL dash characters + collapses whitespace, so a docx
    built from a slightly different markdown revision still matches.

Runs on a COPY. Never touches the original. Prints a full matched/unmatched
report so nothing is applied silently.

    cd backend && venv/Scripts/python scripts/_apply_deai_to_docx.py
Output: docs/report/Dissertation_clean.docx
"""
import re
import subprocess
from pathlib import Path

from docx import Document
import copy

ROOT = Path(r"C:/Users/Bradford Nfor/OneDrive/Desktop/my projects/"
            r"final_year_project")
REPORT = ROOT / "docs" / "report"
SRC = REPORT / "Dissertation.docx"
OUT = REPORT / "Dissertation_clean.docx"
OLD_REV = "b34c947"          # last commit BEFORE the de-AI text edits
FILES = ["front_matter.md", "references.md",
         "chapters/ch1.md", "chapters/ch2.md", "chapters/ch3.md",
         "chapters/ch4.md", "chapters/ch5.md"]

DASHES = "—–‒―"


# ── Markdown -> plain paragraphs ──────────────────────────────────────────────

def strip_inline(s: str) -> str:
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)     # links
    s = s.replace("<br>", "")
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)             # bold
    s = re.sub(r"\*(.+?)\*", r"\1", s)                 # italic
    s = re.sub(r"`(.+?)`", r"\1", s)                   # code
    s = s.replace("\\_", "_").replace("\\*", "*")
    return s.strip()


def md_paragraphs(text: str) -> list[str]:
    """Content paragraphs only: skip headings, tables, and rules. Caption and
    bullet lines are kept (each becomes a docx paragraph)."""
    out = []
    for line in text.splitlines():
        s = line.strip()
        if not s or s.startswith("#") or s.startswith("|"):
            continue
        if s in ("---", "***", "___") or s.startswith("!["):
            continue
        if s.startswith("> "):
            s = s[2:].strip()
        if s.startswith("- ") or s.startswith("* "):
            s = s[2:].strip()
        out.append(strip_inline(s))
    return out


def norm(s: str) -> str:
    for d in DASHES:
        s = s.replace(d, "")
    s = s.replace("-", "")
    return re.sub(r"\s+", " ", s).strip().lower()


def git_show(rev: str, path: str) -> str:
    r = subprocess.run(["git", "show", f"{rev}:{path}"],
                       cwd=ROOT, capture_output=True, text=True, encoding="utf-8")
    return r.stdout


def build_change_map():
    change, collisions = {}, set()
    for rel in FILES:
        old = md_paragraphs(git_show(OLD_REV, f"docs/report/{rel}"))
        new = md_paragraphs((REPORT / rel).read_text(encoding="utf-8"))
        # pair by position where they differ (structure is unchanged: we only
        # rewrote text, never added/removed paragraphs)
        if len(old) != len(new):
            print(f"  !! {rel}: paragraph count {len(old)} -> {len(new)}; "
                  f"pairing by best-effort index")
        for o, n in zip(old, new):
            if o == n:
                continue
            k = norm(o)
            if k in change and change[k] != n:
                collisions.add(k)
            change[k] = n
    for k in collisions:
        change.pop(k, None)
    return change


# ── Formatting-preserving paragraph rebuild ───────────────────────────────────

def formatted_segments(paragraph):
    """Ordered (text, bold, italic, font_name) for runs that carry bold/italic;
    these are the technical terms that must keep their look."""
    segs = []
    for r in paragraph.runs:
        if r.text and (r.bold or r.italic):
            segs.append((r.text, bool(r.bold), bool(r.italic),
                         r.font.name))
    return segs


def has_drawing(paragraph) -> bool:
    """True if any run holds an inline image / drawing, in which case the
    paragraph must never be rebuilt destructively (it would delete the image)."""
    from docx.oxml.ns import qn
    for r in paragraph.runs:
        if r._element.findall(qn("w:drawing")) or \
           r._element.findall(".//" + qn("w:drawing")):
            return True
    return False


def rebuild(paragraph, new_text: str, template_run) -> bool:
    """Rewrite paragraph to new_text, re-inserting its bold/italic segments in
    order. Returns False (and changes nothing) if a segment can't be located in
    new_text, so the caller can report it instead of mangling the paragraph."""
    if has_drawing(paragraph):
        return False        # never destroy a paragraph that carries an image
    segs = formatted_segments(paragraph)
    pieces = []          # list of (text, bold, italic, font)
    pos = 0
    cursor = 0
    for text, b, i, font in segs:
        idx = new_text.find(text, cursor)
        if idx == -1:
            return False        # term not found in new text -> bail, report it
        if idx > cursor:
            pieces.append((new_text[cursor:idx], False, False, None))
        pieces.append((text, b, i, font))
        cursor = idx + len(text)
    if cursor < len(new_text):
        pieces.append((new_text[cursor:], False, False, None))

    # base formatting for plain runs taken from the paragraph's first run
    base_font = template_run.font.name if template_run is not None else None
    base_sz = template_run.font.size if template_run is not None else None

    # clear existing runs
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    for text, b, i, font in pieces:
        run = paragraph.add_run(text)
        run.bold = b or None
        run.italic = i or None
        run.font.name = font or base_font
        if base_sz is not None:
            run.font.size = base_sz
    return True


def main():
    change = build_change_map()
    print(f"changed paragraphs to apply: {len(change)}\n")

    doc = Document(str(SRC))

    def process(paragraphs, where):
        applied = unmatched_fmt = 0
        for p in paragraphs:
            if not p.text.strip():
                continue
            k = norm(p.text)
            if k in change:
                first = next((r for r in p.runs if r.text), None)
                if rebuild(p, change[k], first):
                    applied += 1
                    seen.add(k)
                else:
                    unmatched_fmt += 1
                    report_fmt.append((where, p.text[:70]))
        return applied, unmatched_fmt

    seen = set()
    report_fmt = []
    a1, u1 = process(doc.paragraphs, "body")
    a2 = u2 = 0
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                x, y = process(cell.paragraphs, "table")
                a2 += x; u2 += y

    doc.save(str(OUT))

    print(f"applied in body:   {a1}")
    print(f"applied in tables: {a2}")
    if report_fmt:
        print(f"\nSKIPPED (formatting term not found in new text) "
              f"- handle by hand:")
        for where, txt in report_fmt:
            print(f"  [{where}] {txt}...")
    not_seen = [v for k, v in change.items() if k not in seen]
    if not_seen:
        print(f"\nNOT FOUND in docx ({len(not_seen)}) - likely differ from the "
              f"docx's text revision; verify these by eye:")
        for v in not_seen[:40]:
            print(f"  -> {v[:80]}")
    print(f"\nWrote {OUT.name}")


if __name__ == "__main__":
    main()
