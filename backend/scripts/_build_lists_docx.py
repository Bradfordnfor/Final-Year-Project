r"""Populate the (currently empty) List of Figures and List of Tables in
Dissertation_print.docx.

For each figure/table caption it drops a bookmark, then under the matching
heading it writes one entry per caption: the label + text, a right-aligned dotted
tab leader, and a PAGEREF field pointing at the bookmark. The PAGEREF resolves to
the correct page when the document's fields are updated in WPS/Word (the same
mechanism the main Table of Contents already uses), so no page numbers need to be
known in advance.

Run:  cd backend && venv/Scripts/python scripts/_build_lists_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOC = Path(r"C:/Users/Bradford Nfor/OneDrive/Desktop/my projects/"
           r"final_year_project/docs/report/Dissertation_print.docx")

CAP = re.compile(r"^(Figure|Table)\s*(\d+)\s*:\s*(.*)$")


def add_bookmark(paragraph, name, bid):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def pageref_run(paragraph, bookmark):
    """Append a  {PAGEREF bookmark \h}  field to the paragraph."""
    r = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", f" PAGEREF {bookmark} \\h "),
                       ("sep", None), ("t", "1"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve")
            el.text = text
        elif kind == "t":
            el = OxmlElement("w:t"); el.text = text
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind if kind != "sep" else "separate")
            if kind == "begin":
                el.set(qn("w:fldCharType"), "begin")
            elif kind == "end":
                el.set(qn("w:fldCharType"), "end")
        r._r.append(el)


def make_entry(doc, label_text, bookmark):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(2)
    pf.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = p.add_run(label_text)
    run.font.name = "Times New Roman"; run.font.size = Pt(12)
    p.add_run("\t")
    pageref_run(p, bookmark)
    return p


def main():
    doc = Document(str(DOC))
    paras = doc.paragraphs

    # 1) bookmark every caption, collect (kind, number, text, bookmark)
    figures, tables = [], []
    bid = 5000
    for para in paras:
        m = CAP.match(para.text.strip())
        if not m:
            continue
        kind, num, text = m.group(1), m.group(2), m.group(3).strip().rstrip(".")
        bm = f"_{kind.lower()}_{num}"
        add_bookmark(para, bm, bid); bid += 1
        entry = (f"{kind} {num}:  {text}", bm)
        (figures if kind == "Figure" else tables).append(entry)

    # 2) locate the (empty) heading paragraphs
    def find_heading(title):
        for p in doc.paragraphs:
            if p.style.name.startswith("Heading") and p.text.strip().upper() == title:
                return p
        return None

    hf = find_heading("LIST OF FIGURES")
    ht = find_heading("LIST OF TABLES")

    def insert_after(anchor, entries):
        # build entry paragraphs, then move them right after the anchor heading
        ref = anchor._p
        for label, bm in entries:
            p = make_entry(doc, label, bm)   # created at end of body
            ref.addnext(p._p)                # move after ref
            ref = p._p

    if ht and tables:
        insert_after(ht, tables)
    if hf and figures:
        insert_after(hf, figures)

    doc.save(str(DOC))
    print(f"figures listed: {len(figures)} | tables listed: {len(tables)}")
    print("PAGEREF fields inserted; update fields in WPS to fill page numbers.")


if __name__ == "__main__":
    main()
