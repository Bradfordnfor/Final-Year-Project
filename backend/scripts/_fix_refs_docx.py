r"""Update in-text figure/table references in Dissertation_clean.docx to match
the document's actual sequential labels, and fix a duplicated caption reference.

Mapping (verified from the docx captions):
    Figure 3.N  -> Figure N          (design diagrams, docx Figures 1-11)
    Figure 4.N  -> Figure N+11       (screenshots, docx Figures 12-21)
    Table 4.N   -> Table N           (docx Tables 1-4)

All edits are done at run-text level, never deleting runs, so every image,
caption and formatting run is untouched. Prints a per-change count.
"""
import re
from pathlib import Path
from docx import Document

DOC = Path(r"C:/Users/Bradford Nfor/OneDrive/Desktop/my projects/"
           r"final_year_project/docs/report/Dissertation_clean.docx")

REF = re.compile(r"\b(Figure|Table)\s(\d+)\.(\d+)\b")


def remap(m) -> str:
    kind, a, b = m.group(1), int(m.group(2)), int(m.group(3))
    if kind == "Figure":
        n = b if a == 3 else (b + 11 if a == 4 else None)
    else:                      # Table (only chapter 4 tables exist)
        n = b if a == 4 else None
    if n is None:
        return m.group(0)      # leave anything unexpected untouched
    return f"{kind} {n}"


def main():
    doc = Document(str(DOC))
    counts = {"refs": 0, "dupfix": 0}

    def process(paragraphs):
        for p in paragraphs:
            # 1) fix the duplicated caption reference in the bulk-import paragraph
            for r in p.runs:
                if r.text and "raw error. (Figure 4.2). A course" in r.text:
                    r.text = r.text.replace(
                        "raw error. (Figure 4.2). A course",
                        "raw error. A course")
                    counts["dupfix"] += 1
            # 2) remap figure/table references, run by run
            for r in p.runs:
                if not r.text:
                    continue
                new, n = REF.subn(remap, r.text)
                if n:
                    r.text = new
                    counts["refs"] += n

    process(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                process(cell.paragraphs)

    # verify none of the old N.M style references survive in the body
    leftover = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if re.match(r"^(Figure|Table)\s*\d+\s*:", t):
            continue
        for m in REF.finditer(t):
            if (m.group(1) == "Figure" and m.group(2) in ("3", "4")) or \
               (m.group(1) == "Table" and m.group(2) == "4"):
                leftover.append(m.group(0))

    doc.save(str(DOC))
    print(f"references remapped: {counts['refs']}")
    print(f"duplicate (Figure 4.2) fixed: {counts['dupfix']}")
    print(f"old-style references still in body: {len(leftover)} {sorted(set(leftover))}")


if __name__ == "__main__":
    main()
