"""Assemble the dissertation markdown into a formatted Microsoft Word .docx.

Applies the University of Buea 2026 formatting rules: A4, Times New Roman,
chapter titles 16pt / subheadings 14pt / body 12pt, 1.5 line spacing, justified
body, 2.5cm left-right and 2cm top-bottom margins, Roman page numbers for the
preliminary pages and Arabic from Chapter 1. A Word Table of Contents field is
inserted (right-click -> Update Field in Word). Lists of Tables and Figures are
built from the captions in the chapters.

Run from docs/report/:
    python build_report.py
Output: docs/report/Dissertation.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

HERE = Path(__file__).resolve().parent
CHAPTERS = [HERE / "chapters" / f"ch{i}.md" for i in range(1, 6)]
FRONT_MATTER = HERE / "front_matter.md"
REFERENCES = HERE / "references.md"
OUTPUT = HERE / "Dissertation.docx"

IDENTITY = {
    "country": "REPUBLIC OF CAMEROON",
    "motto": "Peace – Work – Fatherland",
    "university": "UNIVERSITY OF BUEA",
    "faculty": "FACULTY OF ENGINEERING AND TECHNOLOGY",
    "department": "DEPARTMENT OF COMPUTER ENGINEERING",
    "title": "DESIGN AND IMPLEMENTATION OF AN AUTOMATED TIMETABLE "
             "GENERATION AND MANAGEMENT SYSTEM FOR UNIVERSITIES",
    "submission": "A dissertation submitted to the Department of Computer "
                  "Engineering, Faculty of Engineering and Technology, "
                  "University of Buea, in partial fulfilment of the "
                  "requirements for the award of a Bachelor of Engineering "
                  "(B.Eng.) degree in Computer Engineering.",
    "author": "NFOR RINGDAH BRADFORD",
    "matric": "FE22A257",
    "option": "Software Engineering",
    "supervisor": "Dr. Nde Nguti",
    "year": "Academic Year 2025/2026",
}
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


# ── Document styling ──────────────────────────────────────────────────────────

def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)

    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def set_margins(section) -> None:
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)


def set_page_numbering(section, fmt: str, start: int) -> None:
    sectPr = section._sectPr
    for el in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(el)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def add_page_number_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)


def blank_footer(section) -> None:
    section.footer.is_linked_to_previous = False


def add_page_border(section, color: str = "1D4ED8", sz: int = 18) -> None:
    """Blue bordered cover page, per the UB sample."""
    sectPr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "24")
        b.set(qn("w:color"), color)
        borders.append(b)
    sectPr.append(borders)


# ── Inline + block markdown rendering ─────────────────────────────────────────

def add_inline(paragraph, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for i, r in enumerate(cells) if not (i == 1 and set("".join(r)) <= set("-: "))]
    if not cells:
        return
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            if j >= len(table.rows[i].cells):
                continue
            cell = table.rows[i].cells[j]
            cell.paragraphs[0].text = ""
            add_inline(cell.paragraphs[0], val)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(11)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def add_markdown(doc: Document, text: str, chapter_break: bool = True) -> None:
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line.strip() in ("---", "***", "___"):
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_heading(level=1)
            if chapter_break:
                p.paragraph_format.page_break_before = True
            add_inline(p, line[2:].strip())
        elif line.startswith("## "):
            add_inline(doc.add_heading(level=2), line[3:].strip())
        elif line.startswith("### "):
            add_inline(doc.add_heading(level=3), line[4:].strip())
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_inline(p, line[2:].strip())
        elif line.lstrip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line.lstrip()[2:].strip())
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        else:
            add_inline(doc.add_paragraph(), line)
        i += 1


# ── Front matter ──────────────────────────────────────────────────────────────

def parse_sections(md: str) -> dict[str, str]:
    sections, current, buf = {}, None, []
    for line in md.splitlines():
        if line.startswith("## "):
            if current:
                sections[current] = "\n".join(buf).strip()
            current, buf = line[3:].strip(), []
        elif current:
            buf.append(line)
    if current:
        sections[current] = "\n".join(buf).strip()
    return sections


def centered(doc, text, size=12, bold=False, color=None, space=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return p


def build_cover(doc) -> None:
    for t in (IDENTITY["country"], IDENTITY["motto"], IDENTITY["university"],
              IDENTITY["faculty"], IDENTITY["department"]):
        centered(doc, t, size=13, bold=True, color=BLUE)
    centered(doc, "", space=18)
    centered(doc, IDENTITY["title"], size=16, bold=True, color=BLUE, space=18)
    centered(doc, IDENTITY["submission"], size=12, space=18)
    centered(doc, "By", size=12)
    centered(doc, f"{IDENTITY['author']}  ({IDENTITY['matric']})", size=14, bold=True)
    centered(doc, f"Option: {IDENTITY['option']}", size=12, space=18)
    centered(doc, f"Supervisor: {IDENTITY['supervisor']}", size=12, bold=True)
    centered(doc, IDENTITY["year"], size=12)


def build_title_page(doc) -> None:
    for t in (IDENTITY["university"], IDENTITY["faculty"], IDENTITY["department"]):
        centered(doc, t, size=13, bold=True)
    centered(doc, "", space=24)
    centered(doc, IDENTITY["title"], size=16, bold=True, space=24)
    centered(doc, IDENTITY["submission"], size=12, space=18)
    centered(doc, "By", size=12)
    centered(doc, f"{IDENTITY['author']}  ({IDENTITY['matric']})", size=14, bold=True, space=18)
    centered(doc, f"Supervisor: {IDENTITY['supervisor']}", size=12, bold=True)
    centered(doc, IDENTITY["year"], size=12)


def prelim_heading(doc, text):
    p = doc.add_heading(text, level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True


def add_toc_field(doc, instr, placeholder):
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, it, sep, t, end):
        run._r.append(el)


def collect_captions(label: str) -> list[str]:
    found = []
    pat = re.compile(rf">\s*\*\*({label} [\d.]+)\*\*\s*[—:-]\s*(.+)")
    for ch in CHAPTERS:
        for line in ch.read_text(encoding="utf-8").splitlines():
            m = pat.match(line.strip())
            if m:
                cap = m.group(2).split("(")[0].strip().rstrip(".")
                found.append(f"{m.group(1)}: {cap}")
    return found


def build_caption_list(doc, title, label):
    prelim_heading(doc, title)
    items = collect_captions(label)
    if not items:
        doc.add_paragraph("None.")
        return
    for item in items:
        p = doc.add_paragraph(item)
        p.paragraph_format.line_spacing = 1.5


# ── Assembly ──────────────────────────────────────────────────────────────────

def main() -> None:
    doc = Document()
    configure_styles(doc)
    set_margins(doc.sections[0])

    # Section 1 — cover (no page number)
    build_cover(doc)

    # Section 2 — preliminary prose (lower-roman)
    doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(doc.sections[-1])
    fm = parse_sections(FRONT_MATTER.read_text(encoding="utf-8"))

    build_title_page(doc)
    for key in ("Certification", "Dedication", "Acknowledgement", "Abstract"):
        if key in fm:
            prelim_heading(doc, key.upper())
            add_markdown(doc, fm[key], chapter_break=False)

    prelim_heading(doc, "TABLE OF CONTENTS")
    add_toc_field(doc, 'TOC \\o "1-3" \\h \\z \\u',
                  "Right-click and choose Update Field to build the contents.")
    build_caption_list(doc, "LIST OF TABLES", "Table")
    build_caption_list(doc, "LIST OF FIGURES", "Figure")
    if "List of Abbreviations" in fm:
        prelim_heading(doc, "LIST OF ABBREVIATIONS")
        add_markdown(doc, fm["List of Abbreviations"], chapter_break=False)

    # Section 3 — body (decimal, restart at 1)
    doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(doc.sections[-1])
    for ch in CHAPTERS:
        add_markdown(doc, ch.read_text(encoding="utf-8"), chapter_break=True)
    # References as a final unnumbered chapter-level heading
    refs = REFERENCES.read_text(encoding="utf-8")
    refs = re.sub(r"^>.*$", "", refs, flags=re.MULTILINE)
    add_markdown(doc, refs, chapter_break=True)

    # Page-number formats per section
    blank_footer(doc.sections[0])                       # cover: no number
    add_page_border(doc.sections[0])                    # cover: blue border
    set_page_numbering(doc.sections[1], "lowerRoman", 1)
    add_page_number_footer(doc.sections[1])
    set_page_numbering(doc.sections[2], "decimal", 1)
    add_page_number_footer(doc.sections[2])

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
