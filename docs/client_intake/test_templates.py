"""Verify the generated intake templates have the exact expected shape."""
import os
import pytest
from openpyxl import load_workbook

import generate_templates as gt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOMS = os.path.join(HERE, "rooms_template.xlsx")
COURSES = os.path.join(HERE, "courses_template.xlsx")


@pytest.fixture(scope="module", autouse=True)
def _build():
    gt.build_rooms(ROOMS)
    gt.build_courses(COURSES)


def _header(path, sheet):
    wb = load_workbook(path, read_only=True)
    ws = wb[sheet]
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def test_rooms_header_exact():
    assert _header(ROOMS, "Rooms") == [
        "building", "room_name", "capacity", "room_type", "active"]


def test_courses_header_matches_import_format():
    assert _header(COURSES, "Courses") == [
        "code", "name", "level", "department", "semester",
        "weekly_hours", "room_type", "lecturer"]


def test_courses_has_a_shared_department_example():
    wb = load_workbook(COURSES, read_only=True)
    depts = [row[3].value for row in wb["Courses"].iter_rows(min_row=2)]
    assert any(d and "|" in d for d in depts), "need a | shared-course example row"


def test_both_have_instructions_sheet():
    assert "Instructions" in load_workbook(ROOMS, read_only=True).sheetnames
    assert "Instructions" in load_workbook(COURSES, read_only=True).sheetnames
